import io
import docx
from typing import Tuple, Dict, Any

def parse_docx(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """
    Parses a DOCX file from bytes.
    Returns (extracted_text, metadata_dict).
    Raises ValueError on corrupt files or empty text.
    """
    if not file_bytes:
        raise ValueError("Provided DOCX file is empty.")

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(f"Invalid or corrupted DOCX file: {str(e)}")

    extracted_lines = []

    # Process paragraphs
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            extracted_lines.append(text)

    # Process tables
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                extracted_lines.append(" | ".join(row_cells))

    combined_text = "\n".join(extracted_lines).strip()

    if not combined_text or len(combined_text) < 20:
        raise ValueError("We could not extract readable text from this DOCX document.")

    metadata = {
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "character_count": len(combined_text)
    }

    return combined_text, metadata
