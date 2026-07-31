import pytest
import io
import docx
from core.docx_parser import parse_docx

def test_parse_docx_valid():
    doc = docx.Document()
    doc.add_heading("Jane Doe", 0)
    doc.add_paragraph("Email: jane@example.com")
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph("Python, Flask, AWS, Docker, ReactJS")
    
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    text, meta = parse_docx(docx_bytes)
    assert "Jane Doe" in text
    assert "Python" in text
    assert meta["paragraph_count"] >= 3

def test_parse_docx_empty():
    with pytest.raises(ValueError):
        parse_docx(b"")
